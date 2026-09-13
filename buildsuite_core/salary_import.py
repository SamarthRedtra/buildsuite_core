"""Parse and deterministically match the approved 2026 DME salary workbook."""

from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import asdict, dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

COMPANY = "Deltachem Middle East LLC"
EFFECTIVE_DATE = date(2026, 1, 1)
EXPECTED_MISSING = {
	"ahmed asal",
	"alaa sayeed",
	"anees ur rehman",
	"khuloud ali eisa ballekhalib alshamsi",
	"varun venugopal",
}
EXPECTED_TOTALS = {
	"employees": 67,
	"basic": Decimal("237030.00"),
	"hra": Decimal("90600.00"),
	"allowance": Decimal("144325.00"),
	"gross": Decimal("471955.00"),
	"WPS_count": 52,
	"WPS_gross": Decimal("311150.00"),
	"Cash_count": 15,
	"Cash_gross": Decimal("160805.00"),
}
ALIASES = {
	"devaraj gajre pundlikrao": "HR-EMP-00012",
	"adnan rasheed": "HR-EMP-00065",
	"adnan nadeem muhammad nadeem": "HR-EMP-00064",
	"salah uddin moeen uddin": "HR-EMP-00062",
	"mark teakle": "HR-EMP-00052",
	"m a sami": "HR-EMP-00051",
	"moosa ali": "HR-EMP-00067",
	"sohel shaik": "HR-EMP-00068",
}


@dataclass(frozen=True)
class SalaryRow:
	row: int
	mode: str
	name: str
	employee_code: str
	iban: str
	labour_card: str
	basic: Decimal
	hra: Decimal
	allowance: Decimal

	@property
	def gross(self):
		return self.basic + self.hra + self.allowance


@dataclass(frozen=True)
class EmployeeRecord:
	name: str
	employee_name: str
	iban: str
	labour_card: str
	joining_date: date | None


def normalize(value) -> str:
	text = unicodedata.normalize("NFKD", str(value or ""))
	text = "".join(character for character in text if not unicodedata.combining(character))
	return " ".join(re.sub(r"[^a-z0-9]+", " ", text.lower()).split())


def normalize_identifier(value) -> str:
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	normalized = re.sub(r"[^A-Z0-9]", "", str(value).upper())
	return "" if normalized and set(normalized) == {"0"} else normalized


def money(value, row, column) -> Decimal:
	if value in (None, ""):
		return Decimal("0.00")
	try:
		return Decimal(str(value)).quantize(Decimal("0.01"))
	except InvalidOperation as exc:
		raise ValueError(f"Row {row} has an invalid {column} amount: {value!r}") from exc


def parse_workbook(path: str | Path) -> tuple[list[SalaryRow], str]:
	path = Path(path).expanduser().resolve()
	checksum = hashlib.sha256(path.read_bytes()).hexdigest()
	workbook = load_workbook(path, read_only=True, data_only=True)
	try:
		worksheet = workbook.active
		rows = []
		for row_number in range(6, 78):
			values = [worksheet.cell(row_number, column).value for column in range(1, 17)]
			if not values[2]:
				raise ValueError(f"Workbook row {row_number} has no employee name")
			mode = "WPS" if normalize(values[1]) == "wps" else "Cash"
			if normalize(values[1]) not in {"wps", "cash"}:
				raise ValueError(f"Workbook row {row_number} has invalid payment mode {values[1]!r}")
			rows.append(
				SalaryRow(
					row=row_number,
					mode=mode,
					name=str(values[2]).strip(),
					employee_code=normalize_identifier(values[3]),
					iban=normalize_identifier(values[5]),
					basic=money(values[6], row_number, "Basic"),
					hra=money(values[7], row_number, "HRA"),
					allowance=money(values[8], row_number, "Allowance"),
					labour_card=normalize_identifier(values[15]),
				)
			)
	finally:
		workbook.close()
	return rows, checksum


def match_rows(rows: list[SalaryRow], employees: list[EmployeeRecord]):
	indexes = {
		"name": _index(employees, lambda employee: normalize(employee.employee_name)),
		"iban": _index(employees, lambda employee: normalize_identifier(employee.iban)),
		"labour_card": _index(
			employees, lambda employee: normalize_identifier(employee.labour_card)
		),
	}
	matches, missing, errors = [], [], []
	for row in rows:
		candidates = set(indexes["name"].get(normalize(row.name), []))
		if row.iban:
			candidates.update(indexes["iban"].get(row.iban, []))
		if row.labour_card:
			candidates.update(indexes["labour_card"].get(row.labour_card, []))
		if not candidates and normalize(row.name) in ALIASES:
			candidates.add(ALIASES[normalize(row.name)])
		if len(candidates) > 1:
			errors.append(f"Row {row.row} {row.name!r} matched multiple employees: {sorted(candidates)}")
		elif candidates:
			matches.append((row, next(iter(candidates))))
		else:
			missing.append(row)
	return matches, missing, errors


def _index(employees, getter):
	result = {}
	for employee in employees:
		key = getter(employee)
		if key:
			result.setdefault(key, []).append(employee.name)
	return result


def calculate_totals(matches):
	result = {"employees": len(matches)}
	for key in ("basic", "hra", "allowance", "gross"):
		result[key] = sum((getattr(row, key) for row, _ in matches), Decimal("0.00"))
	for mode in ("WPS", "Cash"):
		selected = [row for row, _ in matches if row.mode == mode]
		result[f"{mode}_count"] = len(selected)
		result[f"{mode}_gross"] = sum((row.gross for row in selected), Decimal("0.00"))
	return result


def serialise_row(row):
	return {key: str(value) if isinstance(value, Decimal) else value for key, value in asdict(row).items()}
