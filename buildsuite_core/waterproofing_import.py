"""Strict parser and approved profiles for waterproofing workbooks."""

from __future__ import annotations

import hashlib
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

COMPANY = "Deltachem Middle East LLC"
CUSTOMER_SOURCE = "Desert Group"
CUSTOMER = "Desert Leisure Swimming Pools LLC"
PROJECT_ID = "BS-WP-001"
PROJECT_NAME = "Sobha Hartland II Mansion Lagoon"
EXPECTED_SHA256 = "86f9fa9d96be849e7e30301269ef5d13389c411d358d9f9a39fdf2852f1a66c6"
R2_SHA256_PREFIX = "a76b"
R2_SHA256_SUFFIX = "e643"

EXPECTED = {
	"area": Decimal("318"),
	"thickness": Decimal("3"),
	"man_days": Decimal("3.18"),
	"material": Decimal("46575.43636363637"),
	"labour": Decimal("3466.2"),
	"direct": Decimal("50041.63636363637"),
	"indirect": Decimal("9007.494545454545"),
	"consumables": Decimal("3002.498181818182"),
	"warranty": Decimal("4657.543636363637"),
	"cost": Decimal("66709.17272727273"),
	"selling": Decimal("86496"),
	"profit": Decimal("19786.82727272727"),
}
R2_EXPECTED = {
	"area": Decimal("500"),
	"thickness": Decimal("2"),
	"man_days": Decimal("5"),
	"material": Decimal("23931.81818181818"),
	"labour": Decimal("10900"),
	"direct": Decimal("34831.81818181818"),
	"indirect": Decimal("6269.727272727272"),
	"consumables": Decimal("2089.909090909091"),
	"warranty": Decimal("2393.181818181818"),
	"cost": Decimal("45584.63636363636"),
	"selling": Decimal("86496"),
	"profit": Decimal("40911.36363636364"),
}


@dataclass(frozen=True)
class AssemblyComponent:
	resource: str
	coefficient: Decimal


@dataclass(frozen=True)
class WorkbookProfile:
	key: str
	checksum: str | None
	checksum_prefix: str | None
	checksum_suffix: str | None
	expected: dict[str, Decimal]
	assembly: str
	assembly_name: str
	boq_title: str
	selected: tuple[AssemblyComponent, ...]
	creates_project_context: bool


R1_PROFILE = WorkbookProfile(
	key="waterproofing-3mm-r1",
	checksum=EXPECTED_SHA256,
	checksum_prefix=None,
	checksum_suffix=None,
	expected=EXPECTED,
	assembly="ASM-WB-WATERPROOFING-3MM",
	assembly_name="Workbook Aqualine AL Waterproofing — 3 mm",
	boq_title="Waterproofing Consumption Sheet — Workbook Estimate",
	selected=(
		AssemblyComponent("WB-PRI-SF-PRIMER", Decimal("0.25")),
		AssemblyComponent("WB-FIL-FUMED-SILICA", Decimal("0.2")),
		AssemblyComponent("WB-POL-AQUALINE-AL", Decimal("3") / Decimal("1.1")),
		AssemblyComponent("WB-LAB-SUPERVISOR", Decimal("0.1")),
		AssemblyComponent("WB-LAB-SPRAYER", Decimal("0.1")),
		AssemblyComponent("WB-LAB-WALL-PAINTER", Decimal("0.1")),
		AssemblyComponent("WB-LAB-HELPER", Decimal("0.2")),
	),
	creates_project_context=True,
)
R2_PROFILE = WorkbookProfile(
	key="waterproofing-2mm-r2",
	checksum=None,
	checksum_prefix=R2_SHA256_PREFIX,
	checksum_suffix=R2_SHA256_SUFFIX,
	expected=R2_EXPECTED,
	assembly="ASM-WB-WATERPROOFING-2MM",
	assembly_name="Workbook HP 400 Waterproofing — 2 mm",
	boq_title="Waterproofing Estimate — 500 m² / 2 mm",
	selected=(
		AssemblyComponent("WB-PRI-SF-PRIMER", Decimal("0.25")),
		AssemblyComponent("WB-FIL-AEROSEAL", Decimal("0.2")),
		AssemblyComponent("WB-POL-HP-400", Decimal("1.8181818")),
		AssemblyComponent("WB-LAB-SUPERVISOR", Decimal("0.2")),
		AssemblyComponent("WB-LAB-SPRAYER", Decimal("0.2")),
		AssemblyComponent("WB-LAB-WALL-PAINTER", Decimal("0.2")),
		AssemblyComponent("WB-LAB-HELPER", Decimal("0.4")),
	),
	creates_project_context=False,
)
APPROVED_PROFILES = (R1_PROFILE, R2_PROFILE)


@dataclass(frozen=True)
class RateRow:
	code: str
	name: str
	category: str
	uom: str
	rate: Decimal
	coverage: Decimal
	source_row: int
	source_notes: str


@dataclass(frozen=True)
class WorkbookData:
	checksum: str
	profile: WorkbookProfile
	project_name: str
	customer: str
	location: str
	start_date: date
	prepared_by: str
	revision: str
	values: dict[str, Decimal]
	rates: tuple[RateRow, ...]


def _decimal(value):
	return Decimal(str(value or 0))


def _assert_close(label, actual, expected, tolerance=Decimal("0.01")):
	if abs(actual - expected) > tolerance:
		raise ValueError(f"{label}: expected {expected}, got {actual}")


def _matches_checksum(profile, checksum):
	if profile.checksum:
		return checksum == profile.checksum
	return checksum.startswith(profile.checksum_prefix) and checksum.endswith(profile.checksum_suffix)


def _profile_for_checksum(checksum):
	return next((profile for profile in APPROVED_PROFILES if _matches_checksum(profile, checksum)), None)


def _profile_for_values(values):
	for profile in APPROVED_PROFILES:
		try:
			for key, expected in profile.expected.items():
				_assert_close(key, values[key], expected)
			return profile
		except ValueError:
			continue
	raise ValueError("Workbook totals do not match an approved waterproofing profile")


def _rate_rows(sheet):
	rows = []
	for row in range(7, 11):
		name = str(sheet.cell(row, 2).value).strip()
		rows.append(
			RateRow(
				f"WB-LAB-{name.upper().replace(' ', '-')}",
				name,
				"Labour",
				"Hour",
				_decimal(sheet.cell(row, 3).value),
				_decimal(sheet.cell(row, 4).value),
				row,
				str(sheet.cell(row, 6).value or ""),
			)
		)
	for prefix, category, row_numbers in (
		("PRI", "Material", range(15, 18)),
		("FIL", "Material", range(22, 24)),
		("POL", "Material", range(28, 32)),
	):
		for row in row_numbers:
			name = str(sheet.cell(row, 2).value).strip()
			note = "Coverage m²/unit" if prefix != "POL" else "Workbook kg per m² per 1mm; cached formula uses reciprocal"
			rows.append(
				RateRow(
					f"WB-{prefix}-{name.upper().replace(' ', '-')}",
					name,
					category,
					str(sheet.cell(row, 3).value).strip(),
					_decimal(sheet.cell(row, 4).value),
					_decimal(sheet.cell(row, 5).value),
					row,
					note,
				)
			)
	if len(rows) != 13:
		raise ValueError(f"Expected 13 rate rows, found {len(rows)}")
	return tuple(rows)


def _values(project):
	return {
		"area": _decimal(project["C16"].value),
		"thickness": _decimal(project["C17"].value),
		"man_days": _decimal(project["C18"].value),
		"material": _decimal(project["G25"].value),
		"labour": _decimal(project["G34"].value),
		"direct": _decimal(project["E39"].value),
		"indirect": _decimal(project["F39"].value),
		"consumables": _decimal(project["F40"].value),
		"warranty": _decimal(project["G48"].value),
		"cost": _decimal(project["G49"].value),
		"selling": _decimal(project["G52"].value),
		"profit": _decimal(project["G54"].value),
	}


def parse_workbook(path: str | Path, expected_checksum="approved"):
	path = Path(path).expanduser().resolve()
	checksum = hashlib.sha256(path.read_bytes()).hexdigest()
	if expected_checksum == "approved":
		profile = _profile_for_checksum(checksum)
		if not profile:
			raise ValueError(f"Workbook checksum mismatch: {checksum} is not an approved workbook")
	elif expected_checksum and checksum != expected_checksum:
		raise ValueError(f"Workbook checksum mismatch: expected {expected_checksum}, got {checksum}")
	else:
		profile = None
	workbook = load_workbook(path, read_only=True, data_only=True)
	try:
		if workbook.sheetnames != ["Project Costing", "Material & Rates", "Sheet1"]:
			raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
		project = workbook["Project Costing"]
		values = _values(project)
		profile = profile or _profile_for_values(values)
		for key, expected in profile.expected.items():
			_assert_close(key, values[key], expected)
		start = project["C10"].value
		if isinstance(start, datetime):
			start = start.date()
		data = WorkbookData(
			checksum=checksum,
			profile=profile,
			project_name=str(project["C7"].value or "").strip(),
			customer=str(project["C8"].value or "").strip(),
			location=str(project["C9"].value or "").strip(),
			start_date=start,
			prepared_by=str(project["C11"].value or "").strip(),
			revision=str(project["C12"].value or "").strip(),
			values=values,
			rates=_rate_rows(workbook["Material & Rates"]),
		)
		if data.project_name != PROJECT_NAME or data.customer != CUSTOMER_SOURCE or data.start_date != date(2026, 4, 6):
			raise ValueError("Workbook project identity does not match the approved import")
		return data
	finally:
		workbook.close()


def serialise(data):
	result = asdict(data)
	result["profile"] = data.profile.key
	result["start_date"] = data.start_date.isoformat()
	result["values"] = {key: str(value) for key, value in data.values.items()}
	result["rates"] = [{**asdict(row), "rate": str(row.rate), "coverage": str(row.coverage)} for row in data.rates]
	return result
