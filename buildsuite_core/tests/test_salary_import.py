# Copyright (c) 2026, BuildSuite Core contributors

from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from frappe.tests import UnitTestCase
from openpyxl import Workbook

from buildsuite_core.salary_import import (
	EmployeeRecord,
	SalaryRow,
	match_rows,
	normalize,
	normalize_identifier,
	parse_workbook,
)


class TestSalaryImport(UnitTestCase):
	def test_normalization_and_zero_identifiers(self):
		self.assertEqual(normalize("  José--BIN  "), "jose bin")
		self.assertEqual(normalize_identifier("AE 12-34"), "AE1234")
		self.assertEqual(normalize_identifier(0), "")

	def test_parser_ignores_total_formula_and_uses_components(self):
		with TemporaryDirectory() as directory:
			path = Path(directory) / "salary.xlsx"
			workbook = Workbook()
			worksheet = workbook.active
			for row in range(6, 78):
				worksheet.cell(row, 2, "WPS" if row == 6 else "Cash")
				worksheet.cell(row, 3, f"Employee {row}")
				worksheet.cell(row, 4, row)
				worksheet.cell(row, 7, 100)
				worksheet.cell(row, 8, 20)
				worksheet.cell(row, 9, 30)
				worksheet.cell(row, 11, "=999999")
			workbook.save(path)
			rows, checksum = parse_workbook(path)

		self.assertEqual(len(rows), 72)
		self.assertEqual(rows[0].gross, Decimal("150.00"))
		self.assertEqual(len(checksum), 64)

	def test_matching_is_exact_and_aliases_are_explicit(self):
		employees = [
			EmployeeRecord("HR-EMP-00001", "Exact Person", "AE123", "998", date(2020, 1, 1)),
			EmployeeRecord("HR-EMP-00012", "Different ERP Name", "", "", date(2020, 1, 1)),
		]
		rows = [
			_salary_row(6, "Exact Person", iban="AE123"),
			_salary_row(7, "Devaraj Gajre Pundlikrao"),
			_salary_row(8, "Almost Exact Person"),
		]
		matches, missing, errors = match_rows(rows, employees)

		self.assertEqual([employee for _, employee in matches], ["HR-EMP-00001", "HR-EMP-00012"])
		self.assertEqual([row.row for row in missing], [8])
		self.assertEqual(errors, [])


def _salary_row(row, name, iban=""):
	return SalaryRow(row, "WPS", name, "", iban, "", Decimal("1"), Decimal("2"), Decimal("3"))
