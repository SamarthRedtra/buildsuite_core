from datetime import datetime
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from frappe.tests import UnitTestCase
from openpyxl import Workbook

from buildsuite_core.waterproofing_import import (
	EXPECTED,
	R2_EXPECTED,
	R2_PROFILE,
	parse_workbook,
)
from buildsuite_core.waterproofing_import_documents import SELECTED


class TestWaterproofingImport(UnitTestCase):
	def test_parser_rejects_an_unapproved_checksum(self):
		with TemporaryDirectory() as directory:
			path = Path(directory) / "waterproofing.xlsx"
			_create_workbook(path)
			with self.assertRaisesRegex(ValueError, "checksum mismatch"):
				parse_workbook(path)

	def test_cached_layout_rates_and_totals(self):
		with TemporaryDirectory() as directory:
			path = Path(directory) / "waterproofing.xlsx"
			_create_workbook(path)
			data = parse_workbook(path, expected_checksum=None)

		self.assertEqual(len(data.rates), 13)
		self.assertEqual(data.project_name, "Sobha Hartland II Mansion Lagoon")
		self.assertEqual(data.customer, "Desert Group")
		self.assertEqual(data.values, EXPECTED)
		self.assertEqual(data.start_date.isoformat(), "2026-04-06")

	def test_aqualine_uses_workbook_reciprocal_calculation(self):
		coefficient = dict(SELECTED)["WB-POL-AQUALINE-AL"]
		self.assertEqual(coefficient, Decimal("3") / Decimal("1.1"))
		self.assertNotEqual(coefficient, Decimal("3") * Decimal("1.1"))

	def test_two_mm_profile_accepts_the_approved_fingerprint_and_live_coefficients(self):
		with TemporaryDirectory() as directory:
			path = Path(directory) / "waterproofing-2mm.xlsx"
			_create_workbook(path, R2_EXPECTED)
			with patch(
				"buildsuite_core.waterproofing_import.hashlib.sha256",
				return_value=_Digest("a76b" + "0" * 56 + "e643"),
			):
				data = parse_workbook(path)

		self.assertIs(data.profile, R2_PROFILE)
		self.assertEqual(data.values, R2_EXPECTED)
		self.assertEqual(
			dict((row.resource, row.coefficient) for row in data.profile.selected),
			{
				"WB-PRI-SF-PRIMER": Decimal("0.25"),
				"WB-FIL-AEROSEAL": Decimal("0.2"),
				"WB-POL-HP-400": Decimal("1.8181818"),
				"WB-LAB-SUPERVISOR": Decimal("0.2"),
				"WB-LAB-SPRAYER": Decimal("0.2"),
				"WB-LAB-WALL-PAINTER": Decimal("0.2"),
				"WB-LAB-HELPER": Decimal("0.4"),
			},
		)
		self.assertEqual(data.values["selling"], Decimal("86496"))
		self.assertEqual(
			dict((row.resource, row.coefficient) for row in data.profile.selected)["WB-PRI-SF-PRIMER"]
			* data.values["area"],
			Decimal("125"),
		)
		self.assertAlmostEqual(
			float(dict((row.resource, row.coefficient) for row in data.profile.selected)["WB-POL-HP-400"] * data.values["area"]),
			909.09,
			places=2,
		)


class _Digest:
	def __init__(self, value):
		self.value = value

	def hexdigest(self):
		return self.value


def _create_workbook(path, expected=EXPECTED):
	workbook = Workbook()
	project = workbook.active
	project.title = "Project Costing"
	project["C7"] = "Sobha Hartland II Mansion Lagoon"
	project["C8"] = "Desert Group"
	project["C9"] = "Dubai"
	project["C10"] = datetime(2026, 4, 6)
	for cell, key in (
		("C16", "area"),
		("C17", "thickness"),
		("C18", "man_days"),
		("G25", "material"),
		("G34", "labour"),
		("E39", "direct"),
		("F39", "indirect"),
		("F40", "consumables"),
		("G48", "warranty"),
		("G49", "cost"),
		("G52", "selling"),
		("G54", "profit"),
	):
		project[cell] = float(expected[key])
	rates = workbook.create_sheet("Material & Rates")
	for row, name, rate in (
		(7, "Supervisor", 35),
		(8, "Sprayer", 25),
		(9, "Wall Painter", 17),
		(10, "Helper", 16),
	):
		rates.cell(row, 2, name)
		rates.cell(row, 3, rate)
		rates.cell(row, 4, 10)
		rates.cell(row, 6, "Per person per day")
	for row, name, unit, rate, coverage in (
		(15, "SF Primer", "Litre", 18, 4),
		(16, "SD Primer", "Litre", 25, 4),
		(17, "ME Primer", "Litre", 18, 4),
		(22, "Aeroseal", "Kg", 35, 5),
		(23, "Fumed Silica", "Kg", 28, 5),
		(28, "HP 400", "Kg", 20, 1.1),
		(29, "Aqualine AL", "Kg", 50, 1.1),
		(30, "Hybrid 500", "Kg", 20, 1.1),
		(31, "Standard 300", "Kg", 20, 1.1),
	):
		rates.cell(row, 2, name)
		rates.cell(row, 3, unit)
		rates.cell(row, 4, rate)
		rates.cell(row, 5, coverage)
	workbook.create_sheet("Sheet1")
	workbook.save(path)
